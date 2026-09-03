from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt


OUT = "spotify_memory_engine_analysis_updated.pptx"

BG = RGBColor(18, 18, 18)
PANEL = RGBColor(31, 31, 31)
PANEL_2 = RGBColor(42, 42, 42)
WHITE = RGBColor(245, 245, 245)
MUTED = RGBColor(165, 165, 165)
GREEN = RGBColor(29, 185, 84)
LIME = RGBColor(191, 255, 0)
RED = RGBColor(255, 95, 95)
YELLOW = RGBColor(255, 205, 74)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def no_line(shape):
    shape.line.fill.background()


def box(slide, x, y, w, h, fill=PANEL, radius=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        no_line(shape)
    return shape


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def line(slide, x1, y1, x2, y2, color=GREEN, width=2):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def dot(slide, x, y, r=0.12, color=GREEN):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(r), Inches(r))
    shape.fill.solid(); shape.fill.fore_color.rgb = color; no_line(shape)
    return shape


def base(title, kicker, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = BG
    box(slide, 0, 0, 0.18, 7.5, GREEN)
    text(slide, kicker.upper(), 0.65, 0.42, 5, 0.25, 10, GREEN, True, "Aptos")
    text(slide, title, 0.65, 0.78, 11.8, 0.65, 28, WHITE, True, "Aptos Display")
    text(slide, f"{number:02d}", 12.1, 0.45, 0.55, 0.3, 11, MUTED, True, "Aptos", PP_ALIGN.RIGHT)
    line(slide, 0.65, 1.58, 12.65, 1.58, PANEL_2, 1)
    return slide


def pill(slide, label, x, y, w, color=GREEN, fg=BG):
    box(slide, x, y, w, 0.34, color, True)
    text(slide, label, x, y + 0.07, w, 0.18, 10, fg, True, "Aptos", PP_ALIGN.CENTER)


def bullet(slide, label, x, y, w, accent=GREEN, size=15):
    dot(slide, x, y + 0.08, 0.1, accent)
    text(slide, label, x + 0.22, y, w - 0.22, 0.35, size, WHITE)


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
box(slide, 0, 0, 0.18, 7.5, GREEN)
text(slide, "SPOTIFY AI / BACKEND ANALYSIS", 0.8, 0.75, 7, 0.3, 12, GREEN, True)
text(slide, "Spotify Governed\nAI Memory Engine", 0.8, 1.35, 9.5, 1.6, 38, WHITE, True, "Aptos Display")
text(slide, "A code-grounded walkthrough of personalization, temporal memory, and privacy controls", 0.82, 3.3, 7.6, 0.55, 18, MUTED)
pill(slide, "FASTAPI", 0.82, 4.45, 1.1)
pill(slide, "SQLITE", 2.05, 4.45, 1.05, PANEL_2, WHITE)
pill(slide, "CHROMA", 3.22, 4.45, 1.08, PANEL_2, WHITE)
pill(slide, "NEO4J", 4.42, 4.45, 0.98, PANEL_2, WHITE)
# abstract equalizer / playlist motif
for i, height in enumerate([0.7, 1.5, 1.0, 2.3, 1.25, 2.0, 0.85, 1.7]):
    x = 9.3 + i * 0.38
    box(slide, x, 5.95 - height, 0.16, height, GREEN if i % 3 else LIME, True)
text(slide, "v5.1.0  |  implementation review", 0.82, 6.8, 4, 0.25, 11, MUTED)

# 2. Backend workflow
slide = base("From listening request to personalized result", "backend workflow", 2)
text(slide, "A single chat request passes through validation, context resolution, persistence, and Spotify retrieval.", 0.75, 1.95, 10.5, 0.4, 18, WHITE, True)
workflow = [
    ("01", "Receive", "POST /ai/personalized-chat\nuser_id + message"),
    ("02", "Read", "SQLite profile\ncontrol state"),
    ("03", "Interpret", "Rules resolve mood,\nactivity, artist, bans"),
    ("04", "Remember", "Write profile, history,\ngraph facts, latency"),
    ("05", "Recommend", "Spotify API or\nfallback result"),
]
for i, (n, title, desc) in enumerate(workflow):
    x = 0.75 + i * 2.45
    box(slide, x, 2.85, 2.05, 2.0, PANEL if i != 4 else GREEN, True)
    text(slide, n, x + 0.25, 3.12, 0.38, 0.25, 14, BG if i == 4 else GREEN, True)
    text(slide, title, x + 0.25, 3.55, 1.55, 0.25, 18, BG if i == 4 else WHITE, True, "Aptos Display")
    text(slide, desc, x + 0.25, 4.05, 1.55, 0.55, 13, BG if i == 4 else MUTED)
    if i < len(workflow) - 1:
        line(slide, x + 2.08, 3.85, x + 2.35, 3.85, GREEN, 2)
text(slide, "Control branch: memory_paused skips writes; personalization_enabled removes remembered artist, exclusion, and podcast context.", 0.75, 5.55, 11.5, 0.4, 14, YELLOW)
box(slide, 0.75, 6.2, 11.75, 0.55, PANEL, True)
text(slide, "Returns: ai_dj_response  |  recommended_tracks  |  context_package  |  latency  |  control state", 1.0, 6.38, 11.2, 0.18, 13, WHITE, True, "Aptos", PP_ALIGN.CENTER)

# 3. Folder structure
slide = base("The repository is organized around the memory pipeline", "folder structure", 3)
text(slide, "The checked-in implementation is compact; the README also sketches the larger frontend and pipeline expansion.", 0.75, 1.95, 10.9, 0.4, 17, WHITE, True)
box(slide, 0.75, 2.65, 5.6, 3.55, PANEL, True)
text(slide, "CURRENT BACKEND", 1.1, 2.95, 2.2, 0.22, 11, GREEN, True)
tree = [
    ("backend/", 0, GREEN),
    ("main.py                 API + orchestration", 1, WHITE),
    ("database.py             SQLite helpers", 1, WHITE),
    ("embeddings.py           384-d vector model", 1, WHITE),
    ("vector_db.py            ChromaDB persistence", 1, WHITE),
    ("graph_memory.py         Neo4j adapter", 1, WHITE),
    ("requirements.txt        runtime dependencies", 0, MUTED),
]
for i, (value, indent, color) in enumerate(tree):
    text(slide, ("    " * indent) + value, 1.1, 3.4 + i * 0.38, 4.75, 0.2, 13, color, indent == 0)
box(slide, 6.75, 2.65, 5.75, 3.55, GREEN, True)
text(slide, "README EXPANSION MAP", 7.1, 2.95, 2.8, 0.22, 11, BG, True)
future = [
    "frontend/          client pages + components",
    "memory_pipeline/   auth, ingestion, retrieval",
    "mcp/               governed memory tools",
    "tests/             API, memory, embedding coverage",
    "data/              runtime data mount",
]
for i, value in enumerate(future):
    text(slide, value, 7.1, 3.48 + i * 0.48, 4.8, 0.24, 14, BG)
text(slide, "Implementation read: main.py currently owns the live request path; the standalone vector and Neo4j adapters are ready for integration.", 0.75, 6.55, 11.6, 0.3, 14, YELLOW)

# 4. Contributors
slide = base("Seven contributors, one listening intelligence layer", "project contributors", 4)
text(slide, "A seven-member project team can divide the system into focused ownership areas while sharing one governed user experience.", 0.75, 1.95, 11.1, 0.4, 17, WHITE, True)
contributors = [
    ("01", "Product & UX", "Listening experience\nuser journeys"),
    ("02", "Backend API", "FastAPI routes\nrequest orchestration"),
    ("03", "Memory Systems", "SQLite profile\nhistory + controls"),
    ("04", "AI & Retrieval", "Intent resolution\ncontext package"),
    ("05", "Vector / Graph", "ChromaDB + Neo4j\nsemantic relationships"),
    ("06", "Spotify Integration", "OAuth search\nrecommendation results"),
    ("07", "QA & Deployment", "Tests, secrets\nrelease readiness"),
]
for i, (number, role, detail) in enumerate(contributors):
    col = i % 4
    row = i // 4
    x = 0.75 + col * 3.05
    y = 2.75 + row * 1.55
    fill = GREEN if i == 0 else PANEL
    foreground = BG if i == 0 else WHITE
    accent = BG if i == 0 else GREEN
    box(slide, x, y, 2.65, 1.15, fill, True)
    text(slide, number, x + 0.22, y + 0.2, 0.4, 0.2, 11, accent, True)
    text(slide, role, x + 0.72, y + 0.2, 1.65, 0.23, 15, foreground, True, "Aptos Display")
    text(slide, detail, x + 0.72, y + 0.57, 1.7, 0.36, 12, foreground if i == 0 else MUTED)
box(slide, 0.75, 6.15, 11.75, 0.55, PANEL, True)
text(slide, "Names are intentionally left open because the repository history does not contain a verified seven-member roster.", 1.0, 6.34, 11.2, 0.18, 13, YELLOW, True, "Aptos", PP_ALIGN.CENTER)

# 5. Product intent
slide = base("A personal DJ that remembers with permission", "the product idea", 5)
text(slide, "The backend turns a natural-language listening request into a governed context package and a Spotify search.", 0.7, 1.95, 7.8, 0.55, 20, WHITE, True)
box(slide, 0.7, 2.85, 3.6, 2.6, PANEL, True)
text(slide, "USER SIGNAL", 1.0, 3.15, 2, 0.25, 11, GREEN, True)
text(slide, '"Give me something\nfor a coding session\nwithout rap"', 1.0, 3.55, 2.9, 1.25, 23, WHITE, True, "Aptos Display")
line(slide, 4.55, 4.1, 5.35, 4.1, GREEN, 3)
dot(slide, 5.22, 3.99, 0.22, GREEN)
box(slide, 5.55, 2.85, 3.6, 2.6, PANEL, True)
text(slide, "RESOLVED CONTEXT", 5.85, 3.15, 2.7, 0.25, 11, GREEN, True)
text(slide, "activity: Study Session\nmood: Focus\nbanned: Rap", 5.85, 3.6, 2.8, 1.0, 20, WHITE, True)
line(slide, 9.35, 4.1, 10.15, 4.1, GREEN, 3)
dot(slide, 10.02, 3.99, 0.22, GREEN)
box(slide, 10.35, 2.85, 2.25, 2.6, GREEN, True)
text(slide, "OUTPUT", 10.65, 3.15, 1.5, 0.25, 11, BG, True)
text(slide, "Personalized\ntrack results", 10.65, 3.65, 1.5, 0.8, 21, BG, True, "Aptos Display")
text(slide, "Core promise: relevant recommendations, explainable memory, and user-controlled personalization.", 0.7, 6.25, 11.5, 0.35, 15, MUTED)

# 6. Workflow detail
slide = base("The implemented request path", "runtime workflow", 6)
steps = [
    ("01", "POST /ai/\npersonalized-chat", "FastAPI receives user_id + message"),
    ("02", "Load profile", "SQLite reads memory + control flags"),
    ("03", "Resolve intent", "Keyword rules classify mood, activity, artist, exclusions"),
    ("04", "Persist", "Profile, history, graph facts, and latency are recorded"),
    ("05", "Search Spotify", "OAuth API when configured; deterministic fallback otherwise"),
]
for i, (n, title, desc) in enumerate(steps):
    x = 0.75 + i * 2.45
    dot(slide, x + 0.02, 2.72, 0.32, GREEN)
    text(slide, n, x + 0.02, 2.81, 0.32, 0.14, 9, BG, True, "Aptos", PP_ALIGN.CENTER)
    if i < len(steps) - 1: line(slide, x + 0.35, 2.88, x + 2.22, 2.88, GREEN, 2)
    text(slide, title, x, 3.35, 2.1, 0.6, 18, WHITE, True, "Aptos Display")
    text(slide, desc, x, 4.05, 2.0, 0.8, 13, MUTED)
box(slide, 0.75, 5.6, 11.75, 0.85, PANEL, True)
text(slide, "Response includes", 1.05, 5.86, 1.6, 0.2, 12, GREEN, True)
text(slide, "ai_dj_response   |   recommended_tracks   |   context_package   |   retrieval_latency_ms   |   control state", 2.75, 5.84, 9.2, 0.23, 13, WHITE)

# 7. Architecture
slide = base("One API, three memory representations", "system architecture", 7)
text(slide, "The repository contains a hybrid-memory design: structured state for control, vectors for similarity, and a graph for relationships and lineage.", 0.7, 1.9, 11.4, 0.45, 17, WHITE)
box(slide, 0.75, 2.75, 2.25, 1.65, GREEN, True)
text(slide, "CLIENT", 1.05, 3.08, 1.6, 0.2, 11, BG, True)
text(slide, "HTTP client", 1.05, 3.48, 1.65, 0.55, 18, BG, True, "Aptos Display")
line(slide, 3.15, 3.58, 3.85, 3.58, GREEN, 3)
box(slide, 3.95, 2.75, 2.2, 1.65, PANEL_2, True, GREEN)
text(slide, "ORCHESTRATOR", 4.23, 3.08, 1.65, 0.2, 11, GREEN, True)
text(slide, "FastAPI\nmain.py", 4.23, 3.48, 1.5, 0.55, 19, WHITE, True, "Aptos Display")
line(slide, 6.3, 3.15, 7.0, 3.15, GREEN, 2)
line(slide, 6.3, 3.98, 7.0, 3.98, GREEN, 2)
box(slide, 7.1, 2.25, 2.35, 1.25, PANEL, True)
text(slide, "STRUCTURED STATE", 7.38, 2.55, 1.8, 0.2, 10, GREEN, True)
text(slide, "SQLite profile +\nhistory + controls", 7.38, 2.9, 1.7, 0.45, 16, WHITE, True)
box(slide, 7.1, 3.65, 2.35, 1.25, PANEL, True)
text(slide, "SEMANTIC MEMORY", 7.38, 3.95, 1.8, 0.2, 10, GREEN, True)
text(slide, "ChromaDB +\nMiniLM embeddings", 7.38, 4.3, 1.8, 0.45, 16, WHITE, True)
box(slide, 10.0, 3.0, 2.35, 1.25, PANEL, True)
text(slide, "RELATIONSHIP MEMORY", 10.25, 3.3, 1.85, 0.2, 10, GREEN, True)
text(slide, "Optional Neo4j\ntemporal graph", 10.25, 3.65, 1.7, 0.45, 16, WHITE, True)
line(slide, 9.48, 3.15, 10.0, 3.45, GREEN, 2)
line(slide, 9.48, 4.28, 10.0, 3.8, GREEN, 2)
text(slide, "Repository note: the current `main.py` path directly implements SQLite + a local temporal_graph table. The standalone vector/Neo4j modules are not imported there yet.", 0.75, 5.9, 11.8, 0.45, 13, YELLOW)

# 8. Memory model
slide = base("Memory is a timeline, not a single answer", "data model", 8)
box(slide, 0.75, 2.0, 5.5, 3.8, PANEL, True)
text(slide, "ACTIVE PROFILE", 1.1, 2.35, 2.2, 0.22, 11, GREEN, True)
fields = [("preference_artist", "Diljit Dosanjh"), ("episode_activity", "Study Session"), ("episode_mood", "Focus"), ("exclusion", "Rap"), ("personalization_enabled", "true")]
for i, (key, val) in enumerate(fields):
    y = 2.82 + i * 0.52
    text(slide, key, 1.1, y, 2.7, 0.2, 13, MUTED)
    text(slide, val, 4.1, y, 1.55, 0.2, 13, WHITE, True, "Aptos", PP_ALIGN.RIGHT)
line(slide, 6.75, 3.9, 7.55, 3.9, GREEN, 3)
box(slide, 7.7, 2.0, 4.8, 3.8, PANEL, True)
text(slide, "TEMPORAL GRAPH FACT", 8.05, 2.35, 2.8, 0.22, 11, GREEN, True)
text(slide, "user_42", 8.05, 3.0, 1.3, 0.25, 18, WHITE, True)
text(slide, "PREFERS_ARTIST", 9.55, 3.0, 2.15, 0.25, 14, GREEN, True, "Aptos", PP_ALIGN.CENTER)
text(slide, "Diljit Dosanjh", 10.0, 3.65, 1.8, 0.25, 17, WHITE, True, "Aptos", PP_ALIGN.CENTER)
line(slide, 8.7, 3.12, 10.08, 3.55, GREEN, 2)
text(slide, "confidence 0.98", 8.05, 4.35, 1.8, 0.25, 13, LIME, True)
text(slide, "provenance: User Conversation Intent", 8.05, 4.85, 3.7, 0.25, 12, MUTED)
text(slide, "When a field changes, the previous active fact is closed (`valid_to`) and a new fact becomes active. This creates correction lineage instead of silent overwrites.", 0.75, 6.25, 11.4, 0.35, 14, WHITE)

# 9. Privacy
slide = base("Personalization has visible off-ramps", "trust and control", 9)
text(slide, "The API exposes memory controls as first-class endpoints, not hidden flags.", 0.75, 1.95, 8.2, 0.4, 18, WHITE, True)
controls = [
    ("PAUSE", "Stop profile and history writes during chat", YELLOW),
    ("PERSONALIZATION", "Keep the response generic when disabled", GREEN),
    ("CORRECT", "Update a field and preserve graph provenance", LIME),
    ("PURGE", "Delete SQLite profile, history, and graph facts", RED),
    ("EXPORT", "Return profile, active facts, and telemetry", WHITE),
]
for i, (name, desc, accent) in enumerate(controls):
    y = 2.7 + i * 0.63
    box(slide, 0.8, y, 2.1, 0.42, accent, True)
    text(slide, name, 0.8, y + 0.1, 2.1, 0.18, 11, BG if accent != WHITE else BG, True, "Aptos", PP_ALIGN.CENTER)
    text(slide, desc, 3.25, y + 0.08, 7.9, 0.25, 15, WHITE)
text(slide, "Operational note", 0.8, 6.03, 1.5, 0.2, 12, GREEN, True)
text(slide, "CORS is currently open to all origins, and demo OTPs are returned in the response. Tighten both before production use.", 2.35, 6.01, 10.0, 0.3, 14, YELLOW)

# 10. API surface
slide = base("The backend surface is small and legible", "api map", 10)
groups = [
    ("AUTH", ["POST /auth/send-otp", "POST /auth/register", "POST /auth/login"]),
    ("MEMORY", ["GET /memory/{user_id}", "POST /memory/update-field", "POST /memory/toggle-control", "POST /memory/purge/{user_id}", "GET /memory/export/{user_id}"]),
    ("AI + SEARCH", ["POST /ai/personalized-chat", "GET /ai/retrieve-ranked-memory", "GET /spotify/search"]),
    ("MCP", ["GET /mcp/tools", "search_memory", "correct_memory", "explain_memory_use"]),
]
for i, (name, routes) in enumerate(groups):
    x = 0.75 + (i % 2) * 6.05
    y = 2.05 + (i // 2) * 2.3
    box(slide, x, y, 5.45, 1.85, PANEL, True)
    text(slide, name, x + 0.3, y + 0.25, 2.0, 0.22, 11, GREEN, True)
    for j, route in enumerate(routes):
        text(slide, route, x + 0.3, y + 0.64 + j * 0.24, 4.8, 0.18, 12, WHITE)
text(slide, "Design signal: the MCP manifest already points toward tool-based memory governance, even though the implementation currently exposes the manifest as an API response.", 0.75, 6.75, 11.6, 0.25, 13, MUTED)

# 11. Close / next steps
slide = base("What is working, and what to harden next", "engineering readout", 11)
box(slide, 0.75, 2.0, 5.65, 3.9, PANEL, True)
text(slide, "WORKING FOUNDATION", 1.1, 2.35, 2.7, 0.22, 11, GREEN, True)
for i, value in enumerate(["FastAPI endpoints are easy to discover", "Memory controls cover pause, disable, export, purge", "Temporal facts retain confidence + provenance", "Spotify integration degrades to a local fallback", "Latency is measured and returned per request"]):
    bullet(slide, value, 1.1, 2.85 + i * 0.52, 4.7, GREEN, 14)
box(slide, 6.85, 2.0, 5.65, 3.9, GREEN, True)
text(slide, "NEXT HARDENING MOVES", 7.2, 2.35, 2.9, 0.22, 11, BG, True)
for i, value in enumerate(["Connect vector + Neo4j retrieval to the live chat path", "Replace keyword parsing with a tested intent layer", "Lock down CORS and remove demo OTP leakage", "Add auth, purge, and control-state tests", "Move secrets and DB paths to deployment config"]):
    bullet(slide, value, 7.2, 2.85 + i * 0.52, 4.7, BG, 14)
text(slide, "The strongest next milestone is an end-to-end governed retrieval path: one request, one context package, one auditable recommendation.", 0.75, 6.45, 11.4, 0.35, 16, WHITE, True)

prs.save(OUT)
print(OUT)